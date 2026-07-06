import io
import re
from pathlib import Path
from openpyxl import load_workbook
from pypdf import PdfReader
from app.utils.phone_validator import validate_indian_phone, normalize_phone


def parse_csv(content: bytes) -> list[dict]:
    text = content.decode("utf-8-sig")
    lines = text.strip().split("\n")
    if not lines:
        return []

    headers = [h.strip().lower() for h in lines[0].split(",")]
    rows = []
    for line in lines[1:]:
        values = [v.strip() for v in line.split(",")]
        row = dict(zip(headers, values))
        rows.append(row)
    return rows


def parse_excel(content: bytes) -> list[dict]:
    wb = load_workbook(io.BytesIO(content), read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    headers = [str(h).strip().lower() if h else "" for h in rows[0]]
    result = []
    for row in rows[1:]:
        row_dict = {}
        for i, header in enumerate(headers):
            if header and i < len(row):
                row_dict[header] = str(row[i]).strip() if row[i] else ""
        result.append(row_dict)
    return result


def detect_columns(headers: list[str]) -> dict:
    mapping = {"phone": None, "name": None, "email": None, "company": None}
    phone_aliases = {"phone", "phone_number", "phonenumber", "mobile", "mobile_number", "number", "tel", "telephone", "contact"}
    name_aliases = {"name", "full_name", "fullname", "customer_name", "contact_name", "person_name"}
    email_aliases = {"email", "email_address", "mail", "e-mail"}
    company_aliases = {"company", "organization", "org", "business", "firm", "company_name"}

    for h in headers:
        h_lower = h.lower().strip()
        if h_lower in phone_aliases and mapping["phone"] is None:
            mapping["phone"] = h
        elif h_lower in name_aliases and mapping["name"] is None:
            mapping["name"] = h
        elif h_lower in email_aliases and mapping["email"] is None:
            mapping["email"] = h
        elif h_lower in company_aliases and mapping["company"] is None:
            mapping["company"] = h

    return mapping


def validate_rows(rows: list[dict], column_map: dict) -> tuple[list[dict], list[dict]]:
    valid = []
    errors = []
    seen_phones = set()

    for idx, row in enumerate(rows, start=2):
        row_errors = []
        raw_phone = row.get(column_map.get("phone", ""), "")
        phone = normalize_phone(raw_phone) if raw_phone else ""

        if not raw_phone:
            row_errors.append("Phone number is required")
        elif not validate_indian_phone(raw_phone):
            row_errors.append(f"Invalid Indian phone number: {raw_phone}")

        if phone in seen_phones:
            row_errors.append(f"Duplicate phone number: {phone}")
        if phone:
            seen_phones.add(phone)

        if row_errors:
            errors.append({"row": idx, "phone": raw_phone, "errors": row_errors})
        else:
            valid.append({
                "phone": phone,
                "name": row.get(column_map.get("name", ""), ""),
                "email": row.get(column_map.get("email", ""), ""),
                "company": row.get(column_map.get("company", ""), ""),
            })

    return valid, errors


def parse_pdf(content: bytes) -> list[dict]:
    reader = PdfReader(io.BytesIO(content))
    rows = []
    phone_re = re.compile(r"(\+91)?[6-9]\d{9}")

    for page_num, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        lines = text.strip().split("\n")

        for line in lines:
            line = line.strip()
            if not line:
                continue

            phones = phone_re.findall(line)
            if phones:
                phone = phones[0]
                remaining = line.replace(phone, "").strip()
                parts = [p.strip() for p in re.split(r"[,\t|]+", remaining) if p.strip()]
                rows.append({
                    "phone": phone,
                    "name": parts[0] if len(parts) > 0 else "",
                    "email": parts[1] if len(parts) > 1 else "",
                    "company": parts[2] if len(parts) > 2 else "",
                })

    return rows


def parse_pdf_text(content: bytes) -> str:
    reader = PdfReader(io.BytesIO(content))
    text_parts = []
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text_parts.append(page_text.strip())
    return "\n\n".join(text_parts)


def parse_script_excel(content: bytes) -> list[dict]:
    wb = load_workbook(io.BytesIO(content), read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    headers = [str(h).strip().lower() if h else "" for h in rows[0]]
    name_aliases = {"name", "script_name", "title"}
    content_aliases = {"content", "script", "text", "script_content", "message"}
    lang_aliases = {"language", "lang", "locale"}

    def find_header(aliases):
        for i, h in enumerate(headers):
            if h in aliases:
                return i
        return None

    name_idx = find_header(name_aliases)
    content_idx = find_header(content_aliases)
    lang_idx = find_header(lang_aliases)

    if content_idx is None:
        content_idx = 1 if len(headers) > 1 else 0
    if name_idx is None:
        name_idx = 0 if content_idx != 0 else (1 if len(headers) > 1 else 0)

    result = []
    for row in rows[1:]:
        name_val = str(row[name_idx]).strip() if name_idx is not None and name_idx < len(row) and row[name_idx] else ""
        content_val = str(row[content_idx]).strip() if content_idx is not None and content_idx < len(row) and row[content_idx] else ""
        lang_val = str(row[lang_idx]).strip() if lang_idx is not None and lang_idx < len(row) and row[lang_idx] else "hi-IN"

        if content_val:
            result.append({
                "name": name_val or f"Script {len(result) + 1}",
                "content": content_val,
                "language": lang_val or "hi-IN",
            })

    return result
